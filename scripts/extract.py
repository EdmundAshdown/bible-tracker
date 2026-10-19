"""Regenerate seed.js + bible-data.json from the source spreadsheet.

Usage:
    python3 scripts/extract.py

Reads ' ~/Developer/bible-tracker/ Bible Reading.xlsx' (note leading space),
pulls book metadata + chapter counts from the Tracker sheet and the
translation table from the Translations sheet, and writes seed.js +
bible-data.json.

This is deliberately structure-only: the app starts empty for everyone, so the
spreadsheet's own read/unread checkboxes are counted (how many chapters a book
has) but never exported. To move personal progress between browsers, use the
app's Export / Import buttons.
"""
import json
import os

import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(HERE, " Bible Reading.xlsx")
OUT = os.path.join(HERE, "bible-data.json")


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Tracker"]
    books = []
    for r in range(2, 68):  # rows 2..67 = 66 books
        name = ws.cell(r, 1).value
        count = 0
        for c in range(10, 10 + 160):  # cols J.. = chapter 1..N
            if ws.cell(r, c).value is None:
                break
            count += 1
        books.append(
            {
                "name": name,
                "testament": ws.cell(r, 3).value,
                "classification": ws.cell(r, 4).value,
                "category": ws.cell(r, 5).value,
                "subCategory": ws.cell(r, 6).value,
                "chapters": count,  # how many chapters the book has; no read state
            }
        )

    tr = wb["Translations"]
    translations = []
    for r in range(2, tr.max_row + 1):
        name, abbr, tid = tr.cell(r, 1).value, tr.cell(r, 2).value, tr.cell(r, 3).value
        if name and abbr and tid:
            translations.append({"name": name, "abbr": abbr, "id": int(tid)})

    data = {
        "exportedFrom": " Bible Reading.xlsx",
        "defaultTranslationAbbr": "NKJV",
        "books": books,
        "translations": translations,
    }
    with open(OUT, "w") as f:
        json.dump(data, f, indent=1)
    with open(os.path.join(HERE, "seed.js"), "w") as f:
        f.write("window.BIBLE_SEED = ")
        json.dump(data, f, separators=(",", ":"))
        f.write(";\n")
    total = sum(b["chapters"] for b in books)
    print(f"books={len(books)} chapters={total} translations={len(translations)} -> {OUT}")


if __name__ == "__main__":
    main()
